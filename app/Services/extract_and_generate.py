import sys
import os
import json
import re
import random
import urllib.request

if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

def extract_text_from_file(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    slides_content = []
    
    if ext == '.pdf':
        try:
            import pypdf
            reader = pypdf.PdfReader(file_path)
            for idx, page in enumerate(reader.pages):
                text = page.extract_text() or ''
                raw_lines = [l.strip() for l in text.split('\n') if l.strip()]
                lines = [l for l in raw_lines if len(l) > 1 and not re.match(r'^\d+$', l)]
                title = lines[0] if lines else f'Slide {idx+1}'
                if len(lines) > 1 and len(lines[0]) < 10 and len(lines[1]) > len(lines[0]):
                    title = f"{lines[0]} - {lines[1]}"
                slides_content.append({
                    'slide_num': idx + 1,
                    'title': title,
                    'text': text,
                    'lines': lines
                })
        except Exception as e:
            print(f"Error reading PDF: {e}", file=sys.stderr)
    elif ext in ['.ppt', '.pptx']:
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            for idx, slide in enumerate(prs.slides):
                slide_text = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            t = paragraph.text.strip()
                            if t and len(t) > 1 and not re.match(r'^\d+$', t):
                                slide_text.append(t)
                title = slide_text[0] if slide_text else f'Slide {idx+1}'
                slides_content.append({
                    'slide_num': idx + 1,
                    'title': title,
                    'text': '\n'.join(slide_text),
                    'lines': slide_text
                })
        except Exception as e:
            print(f"Error reading PPT: {e}", file=sys.stderr)
            
    return slides_content

def extract_title_and_subject(filename, slides):
    # 1. Check slide 1 lines
    if slides and len(slides) > 0:
        lines = slides[0].get('lines', [])
        course_code = None
        topic_title = None
        for line in lines:
            m = re.match(r'^(0\d|\d{1,2})\s+(.+)$', line.strip())
            if m:
                week_num = int(m.group(1))
                topic = m.group(2).strip().replace('“', '"').replace('”', '"')
                return f"IF640-W{week_num:02d}: {topic}"
            cm = re.search(r'\b([A-Z]{2,4}\s*\d{3,4})\b', line.strip())
            if cm and not course_code:
                course_code = cm.group(1).replace(' ', '')
            if len(line.strip()) > 8 and not re.search(r'universitas|program studi|fakultas', line.strip(), re.I):
                if not topic_title:
                    topic_title = line.strip()

        if course_code and topic_title:
            return f"{course_code}: {topic_title}"
        elif topic_title:
            return topic_title

    # 2. Check filename pattern like Materi IF640-W06-Social Engineering (1).pdf
    base = os.path.splitext(filename)[0]
    m = re.search(r'IF640[-_\s]*W0?(\d+)[-_\s]*(.+?)(?:_odd\d+|\s*\(\d+\))?$', base, re.IGNORECASE)
    if m:
        week_num = int(m.group(1))
        topic = m.group(2).strip()
        topic = re.sub(r'[-_]', ' ', topic).strip()
        return f"IF640-W{week_num:02d}: {topic}"

    # 3. Clean filename fallback
    if base.startswith('upload_') and slides and len(slides) > 0:
        for line in slides[0].get('lines', []):
            if len(line.strip()) > 5 and not re.search(r'universitas|fakultas|prodi', line.strip(), re.I):
                return line.strip()

    clean = re.sub(r'^Materi\s*', '', base, flags=re.IGNORECASE)
    clean = re.sub(r'_odd\d+', '', clean, flags=re.IGNORECASE)
    clean = re.sub(r'\s*\(\d+\)$', '', clean)
    clean = re.sub(r'[-_]', ' ', clean).strip()
    return clean if clean else filename

def get_env_gemini_key():
    env_file = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), '.env')
    if os.path.exists(env_file):
        try:
            with open(env_file, 'r', encoding='utf-8') as f:
                for line in f:
                    if line.startswith('GEMINI_API_KEY='):
                        val = line.split('=', 1)[1].strip()
                        return val.strip('"').strip("'")
        except Exception:
            pass
    return os.environ.get('GEMINI_API_KEY', '')

def generate_with_gemini(slides, doc_title, gemini_key, is_refresh_quiz=False):
    if not gemini_key:
        return None
    try:
        sample_texts = []
        for s in slides[:40]:
            lines = [x.strip() for x in s.get('lines', []) if len(x.strip()) > 3]
            if lines:
                sample_texts.append(f"[Slide {s['slide_num']}] " + " | ".join(lines[:6]))
        content_context = "\n".join(sample_texts[:30])

        prompt = f"""Kamu adalah Profesor & Dosen Penguji Ujian di universitas terkemuka.
Berikut adalah teks ringkasan materi slide presentasi kuliah (judul awal: '{doc_title}'):
---
{content_context}
---

TUGAS UTAMA:
1. "title": Buat judul materi perkuliahan yang sangat rapi, profesional, dan mencerminkan kode mata kuliah dan topik utama slide jika ada (misal: "AI706: Pengantar Lanskap AI Tools & Penggunaan yang Bertanggung Jawab").
2. "quiz": Buat 4 hingga 5 soal Kuis Pilihan Ganda (Interactive Quiz) yang SANGAT BERKUALITAS TINGGI, menantang, akademis, dan menguji pemahaman konseptual (Active Recall & Critical Thinking).
   - JANGAN membuat soal dangkal atau menaruh judul slide/frasa mentah sebagai pilihan jawaban!
   - Setiap soal wajib memiliki:
     * "question": Pertanyaan jelas dan kontekstual.
     * "options": Array berisi tepat 4 pilihan jawaban yang masuk akal dan gramatikal rapi.
     * "correct": Indeks opsi jawaban yang benar (0, 1, 2, atau 3).
     * "explanation": Penjelasan akademis mendalam mengapa pilihan itu benar dan mengapa pilihan lain keliru.
3. "flashcards": Buat 6 Kartu Flashcard 3D (Active Recall) yang menanyakan konsep esensial pada 'front' dan penjabaran definisi/prinsip lengkap pada 'back' dengan 'badge' kategori.
4. "summary":
   - "executive_summary": Ringkasan eksekutif padat dan komprehensif.
   - "key_takeaways": Array 4-6 poin kunci esensial.
   - "cheat_sheet": Array 4 baris tabel (layer, unit, protocols, hardware).
   - "exam_tips": Array 3 tips strategis menjawab soal ujian materi ini.

Format Output WAJIB JSON murni tanpa markdown:
{{
  "title": "...",
  "quiz": [
    {{
      "question": "...",
      "options": ["A", "B", "C", "D"],
      "correct": 0,
      "explanation": "..."
    }}
  ],
  "flashcards": [
    {{
      "front": "...",
      "back": "...",
      "badge": "..."
    }}
  ],
  "summary": {{
    "executive_summary": "...",
    "key_takeaways": ["..."],
    "cheat_sheet": [
      {{"layer": "...", "unit": "...", "protocols": "...", "hardware": "..."}}
    ],
    "exam_tips": ["..."]
  }}
}}
"""
        models = [
            'gemini-3.1-flash-lite-preview',
            'gemini-3-flash-preview',
            'gemini-2.5-flash',
            'gemini-flash-latest'
        ]
        for m in models:
            url = f"https://generativelanguage.googleapis.com/v1beta/models/{m}:generateContent?key={gemini_key}"
            payload = {
                'contents': [{'parts': [{'text': prompt}]}],
                'generationConfig': {
                    'temperature': 0.7 if is_refresh_quiz else 0.3,
                    'maxOutputTokens': 3500
                }
            }
            req = urllib.request.Request(
                url,
                data=json.dumps(payload).encode('utf-8'),
                headers={'Content-Type': 'application/json'}
            )
            try:
                with urllib.request.urlopen(req, timeout=25) as resp:
                    if resp.status == 200:
                        raw = json.loads(resp.read().decode('utf-8'))
                        if 'candidates' in raw and raw['candidates']:
                            txt = raw['candidates'][0]['content']['parts'][0]['text'].strip()
                            if txt.startswith('```json'):
                                txt = txt[7:]
                            if txt.startswith('```'):
                                txt = txt[3:]
                            if txt.endswith('```'):
                                txt = txt[:-3]
                            txt = txt.strip()
                            data = json.loads(txt)
                            if data.get('quiz') and len(data['quiz']) >= 3:
                                return data
            except Exception:
                continue
    except Exception:
        pass
    return None

def shuffle_quiz_options(quiz_list):
    result = []
    for item in quiz_list:
        opts = list(item['options'])
        correct_text = opts[item['correct']]
        random.shuffle(opts)
        new_correct = opts.index(correct_text)
        result.append({
            'question': item['question'],
            'options': opts,
            'correct': new_correct,
            'explanation': item['explanation']
        })
    return result

def analyze_and_generate(file_path, slides, gemini_key='', orig_title='', orig_filename='', is_refresh_quiz=False):
    filename = orig_filename if orig_filename else os.path.basename(file_path)
    combined_text = ' '.join([s['text'] for s in slides]).lower()
    total_slides = len(slides) if slides else 20
    doc_title = orig_title if orig_title and not orig_title.startswith('upload_') else extract_title_and_subject(filename, slides)
    dt_lower = doc_title.lower()
    fn_lower = filename.lower()
    
    # 0. Try Gemini API first for any presentation (returns world-class custom quiz & flashcards)
    if not gemini_key:
        gemini_key = get_env_gemini_key()
    
    if gemini_key:
        gemini_result = generate_with_gemini(slides, doc_title, gemini_key, is_refresh_quiz=is_refresh_quiz)
        if gemini_result:
            title_final = gemini_result.get('title') or doc_title
            # Ensure video structure is present
            video_data = {
                'duration': '4:15',
                'lecturer': 'Dr. Maya AI - Virtual Lecturer',
                'chapters': [
                    {'time': '00:00', 'title': f'Pengantar {title_final[:30]}'},
                    {'time': '01:20', 'title': 'Lanskap & Konsep Utama'},
                    {'time': '02:45', 'title': 'Etika & Evaluasi Kritis'},
                    {'time': '03:40', 'title': 'Rangkuman Ujian'}
                ],
                'slides': [
                    {
                        'slide_num': 1,
                        'title': title_final,
                        'subtitle': 'Panduan Belajar Cerdas AI',
                        'bullet_points': gemini_result.get('summary', {}).get('key_takeaways', ['Pahami konsep inti', 'Evaluasi implementasi'])[:3],
                        'transcript': f'Halo rekan mahasiswa! Di modul ini kita akan membedah secara mendalam topik {title_final}.'
                    }
                ]
            }
            return {
                'title': title_final,
                'slides_count': total_slides,
                'summary': gemini_result.get('summary', {}),
                'quiz': gemini_result.get('quiz', []),
                'flashcards': gemini_result.get('flashcards', []),
                'video': video_data
            }
    
    # 1. Week 05: Physics of Malware
    if 'w05' in dt_lower or 'physics of malware' in dt_lower or 'w05' in fn_lower or 'physics of malware' in fn_lower:
        return generate_malware_w05(doc_title, total_slides)
        
    # 2. Week 04: Unsupervised Learning & Clustering
    if 'w04' in dt_lower or 'unsupervised' in dt_lower or 'pattern recognition' in dt_lower or 'w04' in fn_lower:
        return generate_unsupervised_w04(doc_title, total_slides)
        
    # 3. Week 03: Supervised Learning & Threat Classification
    if 'w03' in dt_lower or 'supervised' in dt_lower or 'classify threats' in dt_lower or 'w03' in fn_lower:
        return generate_supervised_w03(doc_title, total_slides)
        
    # 4. Week 02: How AI Thinks about Data
    if 'w02' in dt_lower or 'thinks' in dt_lower or 'w02' in fn_lower:
        return generate_data_w02(doc_title, total_slides)
        
    # 5. Week 01: The AI Cyber Landscape
    if 'w01' in dt_lower or 'cyber landscape' in dt_lower or 'w01' in fn_lower:
        return generate_landscape_w01(doc_title, total_slides)
        
    # 6. Week 06: Social Engineering
    if 'w06' in dt_lower or 'social engineering' in dt_lower or 'w06' in fn_lower or 'social engineering' in fn_lower:
        return generate_social_engineering_w06(doc_title, total_slides)
        return generate_social_engineering_w06(doc_title, total_slides)
        
    # Adaptive extraction for any other presentation
    return generate_adaptive_data(doc_title, filename, slides, combined_text)

# ==============================================================================
# WEEK 04: PATTERN RECOGNITION ATTACK WITH UNSUPERVISED LEARNING
# ==============================================================================
def generate_unsupervised_w04(doc_title, total_slides):
    quiz_pool = [
        {
            'question': 'Mengapa algoritma K-Means dan algoritma berbasis jarak mewajibkan normalisasi fitur (StandardScaler / MinMaxScaler) sebelum menghitung jarak Euclidean?',
            'options': [
                'Karena fitur dengan rentang nilai besar (misal: Bytes Transferred jutaan) akan mendominasi perhitungan jarak dan mengabaikan fitur bernilai kecil (misal: Port atau Flag count)',
                'Karena model K-Means tidak dapat membaca angka positif jika belum distandarisasi',
                'Karena normalisasi mengubah data menjadi file audio berformat WAV',
                'Karena sistem operasi Windows memblokir data mentah yang tidak dienkripsi'
            ],
            'correct': 0,
            'explanation': 'Perhitungan jarak Euclidean d(x,y) = sqrt(Σ(xi - yi)^2) sangat sensitif terhadap skala fitur. Tanpa standarisasi (z-score), fitur berangka besar akan mendikte hasil cluster dan melenyapkan kontribusi fitur penting lainnya.'
        },
        {
            'question': 'Bagaimana cara kerja Elbow Method dalam menentukan jumlah cluster K yang optimal pada K-Means?',
            'options': [
                'Mencari titik sudut (siku) pada grafik nilai Inertia/SSE di mana penambahan nilai K tidak lagi menurunkan error secara signifikan',
                'Memilih nilai K yang menghasilkan jumlah data ganjil pada setiap cluster',
                'Menghitung waktu proses CPU hingga mencapai 100%',
                'Mengalikan jumlah baris dataset dengan akar kuadrat jumlah fitur'
            ],
            'correct': 0,
            'explanation': 'Elbow Method memplot nilai Within-Cluster Sum of Squares (Inertia) terhadap variasi nilai K. Titik belok/siku ("elbow") menunjukkan batas di mana penambahan cluster baru hanya memberikan perbaikan marginal.'
        },
        {
            'question': 'Apa indikasi jika Silhouette Score hasil clustering bernilai tinggi (mendekati +1.0)?',
            'options': [
                'Sampel berada sangat dekat dengan pusat cluster-nya sendiri (kohesi tinggi) dan terpisah jauh dari cluster tetangga (separasi baik)',
                'Model mengalami overfitting parah dan harus dibatalkan',
                'Semua data berhasil dilabeli sebagai malware ransomware',
                'Dataset tidak memiliki nilai numerik sama sekali'
            ],
            'correct': 0,
            'explanation': 'Silhouette Score mengukur (b - a) / max(a, b). Nilai mendekati +1 menunjukkan titik data sangat padat di dalam cluster-nya dan terpisah tegas dari cluster lainnya. Nilai < 0 menunjukkan salah pengelompokan.'
        },
        {
            'question': 'Manakah keunggulan utama algoritma DBSCAN (Density-Based Spatial Clustering) jika dibandingkan dengan K-Means dalam deteksi serangan siber?',
            'options': [
                'DBSCAN mampu menemukan cluster dengan bentuk arbitrer (non-spherical) dan secara otomatis menandai outlier berdensitas rendah sebagai noise (-1)',
                'DBSCAN tidak membutuhkan memori RAM sama sekali',
                'DBSCAN selalu menghasilkan tepat 2 cluster',
                'DBSCAN hanya dapat memproses data gambar JPEG'
            ],
            'correct': 0,
            'explanation': 'K-Means mengasumsikan cluster berbentuk bola dan memaksa setiap outlier masuk ke salah satu cluster. DBSCAN mengelompokkan berdasarkan kepadatan tetangga (Epsilon & min_samples) dan memisahkan anomali sebagai noise (-1).'
        },
        {
            'question': 'Dalam analisis keamanan siber menggunakan Hierarchical Clustering, apa fungsi utama dari visualisasi Dendrogram?',
            'options': [
                'Menyediakan jejak audit visual (audit trail) tentang hierarki penggabungan dan tingkat kemiripan behavioral antar sampel',
                'Mempercepat kecepatan koneksi internet firewall',
                'Menghapus log event Windows yang mencurigakan',
                'Mengekstrak password administrator dari memori RAM'
            ],
            'correct': 0,
            'explanation': 'Dendrogram memvisualisasikan bagaimana cluster bergabung secara bertingkat (agglomerative). Analis dapat memotong pohon dendrogram pada jarak tertentu untuk melacak silsilah kemiripan rumpun varian malware.'
        },
        {
            'question': 'Pada algoritma DBSCAN, apa syarat yang harus dipenuhi agar suatu titik data diklasifikasikan sebagai Core Point?',
            'options': [
                'Memiliki minimal sejumlah `min_samples` titik tetangga di dalam radius jarak `Epsilon (ε)`',
                'Berada tepat di koordinat (0, 0)',
                'Memiliki label ground truth bertipe binary',
                'Telah diproses menggunakan algoritma Decision Tree'
            ],
            'correct': 0,
            'explanation': 'Core Point adalah titik yang memiliki minimal `min_samples` tetangga dalam radius ε. Jika kurang tetapi berada di dekat Core Point, ia menjadi Border Point; jika terisolasi, ia menjadi Noise Point.'
        },
        {
            'question': 'Mengapa visualisasi 2D menggunakan PCA (Principal Component Analysis) tidak boleh dijadikan satu-satunya tolok ukur dalam menyimpulkan kualitas clustering?',
            'options': [
                'Karena PCA memproyeksikan fitur berdimensi tinggi ke 2 dimensi, sehingga distorsi jarak dan tumpang tindih visual dapat terjadi meskipun cluster terpisah baik di ruang aslinya',
                'Karena PCA menghapus seluruh nilai numerik dataset',
                'Karena PCA hanya berlaku untuk data teks tanpa angka',
                'Karena PCA membutuhkan lisensi berbayar dari Microsoft'
            ],
            'correct': 0,
            'explanation': 'PCA adalah alat proyeksi linear untuk visualisasi manusia. Reduksi dimensi ke 2D membuang sebagian varians data, sehingga cluster yang tampak tumpang tindih di visualisasi 2D bisa jadi sebenarnya terpisah sempurna di ruang multidimensi.'
        },
        {
            'question': 'Bagaimana cara mendeteksi serangan anomali pada lalu lintas jaringan menggunakan model K-Means yang telah dilatih pada data normal?',
            'options': [
                'Menghitung jarak titik data baru ke centroid terdekat; jika jarak melebihi ambang batas (threshold), data ditandai sebagai anomali',
                'Mengganti nama file log menjadi format ZIP',
                'Mematikan semua port koneksi TCP/IP',
                'Menghitung jumlah karakter nama domain URL'
            ],
            'correct': 0,
            'explanation': 'Pada K-Means Anomaly Detection, data jaringan normal berada di dekat centroid. Paket serangan baru yang tidak wajar akan menghasilkan jarak yang sangat jauh dari semua centroid normal (Distance-to-Centroid Anomaly Scoring).'
        }
    ]

    selected_quiz = random.sample(quiz_pool, min(6, len(quiz_pool)))
    shuffled_quiz = shuffle_quiz_options(selected_quiz)

    return {
        'title': doc_title if doc_title else 'IF640-W04: Pattern Recognition Attack with Unsupervised Learning',
        'slides_count': total_slides,
        'summary': {
            'executive_summary': (
                f"Materi kuliah '{doc_title}' ({total_slides} slides) mengkaji penerapan teknik Unsupervised Machine Learning "
                "dalam domain keamanan siber untuk mendeteksi ancaman tanpa bergantung pada label historis (unlabeled data). "
                "Fokus materi mencakup pemetaan telemetri mentah ke ruang fitur (feature space), urgensi normalisasi fitur "
                "(StandardScaler & MinMaxScaler) terhadap pengukuran jarak (Euclidean, Manhattan, Cosine), "
                "pemodelan pengelompokan berbasis centroid via K-Means (seleksi K optimal dengan Elbow Method & Silhouette Score), "
                "deteksi outlier berbasis densitas via DBSCAN (Core, Border, Noise Points), serta Hierarchical Clustering "
                "(Dendrogram & Linkage Criteria) untuk analisis kemiripan varian malware. "
                "Pembahasan diakhiri dengan evaluasi model tanpa label (Davies-Bouldin, Silhouette) dan proyeksi dimensi PCA."
            ),
            'key_takeaways': [
                'Clustering vs Anomaly Detection: Clustering mengelompokkan pola aktivitas jaringan normal ke dalam baseline cluster, sedangkan Anomaly Detection mengidentifikasi deviasi titik data pencilan (outliers) yang berada jauh dari batas densitas baseline.',
                'Pentingnya Normalisasi Fitur (Scaling Matters): Metrik jarak (seperti Euclidean Distance) sangat sensitif terhadap skala fitur. Fitur dengan rentang nilai besar (misal: Bytes Transferred ribuan s/d jutaan) akan mendominasi fitur berrentang kecil (misal: Port Number atau Flag count) jika tidak distandarisasi menggunakan StandardScaler (z-score) atau MinMaxScaler.',
                'K-Means Clustering & Seleksi K: Algoritma partisi berbasis centroid. Menentukan nilai K optimal dilakukan melalui Elbow Method (mencari titik siku pada kurva Within-Cluster SSE/Inertia) dan Silhouette Score (mengukur kohesi internal vs separasi eksternal pada rentang -1 hingga +1). K-Means sensitif terhadap outlier dan mengasumsikan cluster berbentuk bola (spherical).',
                'DBSCAN (Density-Based Spatial Clustering): Mampu menemukan cluster dengan bentuk sembarang (arbitrary shape) dan secara otomatis menandai outlier berdensitas rendah sebagai noise (-1). Dikonfigurasi melalui dua parameter inti: radius Epsilon (ε) dan min_samples untuk mengklasifikasi titik menjadi Core, Border, atau Noise Points.',
                'Hierarchical Clustering & Dendrogram: Pendekatan Agglomerative (bottom-up) menggabungkan titik data secara bertingkat berdasarkan kriteria Linkage (Single, Complete, Average, Ward) hingga membentuk pohon dendrogram. Sangat efektif sebagai jejak audit (audit trail) kemiripan behavioral rumpun malware tanpa perlu menentukan K di awal.',
                'Evaluasi Tanpa Label & Visualisasi PCA: Kualitas model tanpa ground-truth dievaluasi menggunakan metrik internal (Silhouette Coefficient, Davies-Bouldin Index). PCA digunakan untuk memproyeksikan fitur multi-dimensi ke 2D demi kemudahan visualisasi analis SOC, namun PCA bukan model clustering itu sendiri.'
            ],
            'cheat_sheet': [
                {
                    'layer': 'Distance & Feature Scaling',
                    'unit': 'Normalisasi Metrik',
                    'protocols': 'Euclidean, Manhattan, Cosine | StandardScaler: (x - mean) / std',
                    'hardware': 'Mencegah fitur angka besar (bytes) mendominasi ruang metrik'
                },
                {
                    'layer': 'K-Means Clustering',
                    'unit': 'Centroid Partition',
                    'protocols': 'Assignment step -> Update centroid -> Hitung Inertia SSE',
                    'hardware': 'Grouping netflow normal; Anomaly score = distance to centroid'
                },
                {
                    'layer': 'Optimal K Selection',
                    'unit': 'Validasi Model',
                    'protocols': 'Elbow Method (Inertia SSE) & Silhouette Score (-1 s/d +1)',
                    'hardware': 'K optimal saat Silhouette > 0.5 dan kurva SSE melandai'
                },
                {
                    'layer': 'DBSCAN Density Model',
                    'unit': 'Non-Parametric K',
                    'protocols': 'Radius Epsilon (ε) & min_samples (Core, Border, Noise)',
                    'hardware': 'Mendeteksi cluster bentuk bebas & isolasi serangan sebagai noise (-1)'
                },
                {
                    'layer': 'Hierarchical / Tree',
                    'unit': 'Agglomerative',
                    'protocols': 'Linkage: Single (min), Complete (max), Ward (variance)',
                    'hardware': 'Dendrogram audit trail taksonomi kemiripan malware family'
                },
                {
                    'layer': 'Evaluasi & Visualisasi',
                    'unit': 'Internal Scoring',
                    'protocols': 'Silhouette Index, Davies-Bouldin, PCA 2D/3D projection',
                    'hardware': 'Validasi kepadatan cluster SOC tanpa label ground-truth'
                }
            ],
            'exam_tips': [
                'Jebakan Ujian Efek Skala Fitur: Jika soal menanyakan dampak menghitung jarak Euclidean pada fitur durasi (0-60 detik) dan ukuran paket (0-1.500.000 bytes) tanpa normalisasi, jawabannya: Fitur ukuran paket akan mendominasi 99.9% hasil jarak, membuat fitur durasi tidak berdampak.',
                'Kelemahan K-Means vs Keunggulan DBSCAN: K-Means gagal mendeteksi cluster yang memanjang atau melingkar (non-spherical) dan mewajibkan kita menentukan nilai K di awal. DBSCAN unggul karena tidak memerlukan input jumlah K dan secara otomatis melabeli data terpencil sebagai noise (-1).',
                'Interpretasi Nilai Silhouette Score: Nilai mendekati +1.0 berarti sampel sangat dekat dengan cluster-nya sendiri dan jauh dari cluster tetangga. Nilai 0.0 berada di perbatasan, sedangkan nilai negatif mengindikasikan kemungkinan salah cluster.',
                'Kriteria Linkage Dendrogram: Ward linkage meminimalkan kenaikan total varians dalam cluster saat menggabungkan dua cluster, menghasilkan ukuran cluster yang paling seimbang dan kompak.'
            ]
        },
        'quiz': shuffled_quiz,
        'flashcards': [
            {
                'front': 'StandardScaler (Feature Normalization)',
                'back': 'Rumus: (x - mean) / std. Menyamakan skala seluruh fitur ke mean 0 dan varians 1 agar perhitungan jarak Euclidean tidak didominasi oleh fitur dengan angka besar.',
                'badge': 'Feature Engineering'
            },
            {
                'front': 'Elbow Method pada K-Means',
                'back': 'Teknik visualisasi kurva Sum of Squared Errors (Inertia) terhadap nilai K. Titik sudut siku ("elbow") menandakan penambahan K baru hanya menurunkan error secara marginal.',
                'badge': 'Optimasi K-Means'
            },
            {
                'front': 'Silhouette Score (-1 s/d +1)',
                'back': 'Metrik kohesi intra-cluster vs separasi inter-cluster. Nilai mendekati +1 menandakan cluster padat dan terpisah baik; nilai < 0 mengindikasikan sampel salah cluster.',
                'badge': 'Evaluasi Model'
            },
            {
                'front': 'DBSCAN Core vs Noise Point',
                'back': 'Core: memiliki minimal min_samples titik dalam radius ε. Noise: titik terpencil yang tidak masuk core atau border point (dilabeli -1 sebagai anomali).',
                'badge': 'Density Clustering'
            },
            {
                'front': 'Dendrogram pada Hierarchical Clustering',
                'back': 'Diagram pohon hierarki yang memvisualisasikan urutan penggabungan cluster pada Agglomerative Clustering dari level individu hingga satu kesatuan.',
                'badge': 'Hierarchical Clustering'
            },
            {
                'front': 'Distance-to-Centroid Anomaly Scoring',
                'back': 'Metode deteksi ancaman di mana skor keanehan suatu event jaringan dihitung dari seberapa jauh jaraknya ke titik pusat (centroid) cluster normal.',
                'badge': 'Deteksi Anomali'
            },
            {
                'front': 'K-Means vs DBSCAN',
                'back': 'K-Means: Memerlukan parameter K, mengasumsikan cluster bulat (spherical), sensitif outlier. DBSCAN: Menemukan K otomatis, mendeteksi cluster bentuk bebas, mengisolasi noise.',
                'badge': 'Perbandingan Algoritma'
            },
            {
                'front': 'PCA (Principal Component Analysis)',
                'back': 'Teknik reduksi dimensi linear untuk memproyeksikan fitur berdimensi tinggi ke 2D/3D demi kemudahan visualisasi analis, bukan model clustering itu sendiri.',
                'badge': 'Visualisasi Dimensi'
            }
        ],
        'video': {
            'duration': '4:20',
            'lecturer': 'Dr. Maya AI - Virtual Lecturer',
            'chapters': [
                {'time': '00:00', 'title': 'Bab 1: Pengantar Unsupervised Learning dalam Keamanan Siber'},
                {'time': '01:05', 'title': 'Bab 2: K-Means Clustering, Normalisasi Jarak, & Elbow Method'},
                {'time': '02:15', 'title': 'Bab 3: DBSCAN & Deteksi Anomali Berbasis Kepadatan'},
                {'time': '03:25', 'title': 'Bab 4: Hierarchical Clustering, Dendrogram, & Evaluasi SOC'}
            ],
            'slides': [
                {
                    'slide_num': 1,
                    'title': 'Unsupervised Learning dalam Keamanan Siber',
                    'subtitle': 'Menemukan Pola Serangan Tersembunyi Tanpa Label Historis',
                    'bullet_points': [
                        'Menjawab tantangan ancaman Zero-Day yang belum memiliki signature biner.',
                        'Mengelompokkan event log jaringan menjadi profil baseline aktivitas sah.',
                        'Mendeteksi anomali sebagai deviasi signifikan dari klaster normal.'
                    ],
                    'transcript': 'Selamat datang di kuliah Unsupervised Learning in Cybersecurity! Di modul ini, kita mempelajari bagaimana algoritma AI dapat menemukan struktur dan pola ancaman tersembunyi dari miliaran event log tanpa memerlukan label data historis.'
                },
                {
                    'slide_num': 2,
                    'title': 'K-Means, Jarak Euclidean, & Seleksi K',
                    'subtitle': 'Standarisasi Fitur dan Optimasi Elbow Method',
                    'bullet_points': [
                        'Normalisasi fitur (StandardScaler) wajib dilakukan agar fitur byte tidak mendominasi.',
                        'Elbow Method dan Silhouette Score menentukan jumlah cluster K paling optimal.',
                        'Jarak data baru ke centroid terdekat dijadikan skor indikator anomali.'
                    ],
                    'transcript': 'K-Means bekerja dengan mengiterasi posisi centroid. Ingatlah bahwa standarisasi fitur sangat krusial, karena jarak Euclidean akan bias bila ada perbedaan skala yang ekstrem antar fitur telemetri jaringan Anda.'
                },
                {
                    'slide_num': 3,
                    'title': 'DBSCAN: Density-Based Clustering',
                    'subtitle': 'Menemukan Klaster Bentuk Bebas & Mengisolasi Noise (-1)',
                    'bullet_points': [
                        'Parameter kunci: Radius Epsilon (ε) dan batas minimum tetangga (min_samples).',
                        'Mampu mengenali serangan anomali terpencil dan mengelompokkannya sebagai noise.',
                        'Tidak mengasumsikan klaster berbentuk bulat seperti halnya K-Means.'
                    ],
                    'transcript': 'Berbeda dengan K-Means, DBSCAN tidak memaksa setiap data masuk ke dalam klaster. Titik yang berada di daerah berdensitas rendah akan langsung dilabeli sebagai noise bernilai minus satu, yang sangat ideal untuk deteksi serangan siber.'
                },
                {
                    'slide_num': 4,
                    'title': 'Hierarchical Dendrogram & Evaluasi',
                    'subtitle': 'Audit Kemiripan Varian Malware dan Validasi Tanpa Label',
                    'bullet_points': [
                        'Dendrogram memvisualisasikan tingkatan kesamaan perilaku rumpun malware.',
                        'Metrik internal seperti Silhouette dan Davies-Bouldin mengukur kepadatan klaster.',
                        'PCA digunakan untuk proyeksi visual 2 dimensi, namun bukan pengganti model klaster.'
                    ],
                    'transcript': 'Sebagai penutup, Hierarchical Clustering memberikan jejak audit visual melalui dendrogram untuk mengamati kekerabatan varian malware. Gunakan metrik Silhouette Score untuk memvalidasi kepadatan klaster Anda sebelum diimplementasikan di SOC.'
                }
            ]
        }
    }

# ==============================================================================
# WEEK 05: THE PHYSICS OF MALWARE
# ==============================================================================
def generate_malware_w05(doc_title, total_slides):
    quiz_pool = [
        {
            'question': 'Berdasarkan materi PE Header, apa indikasi kecurigaan utama jika Virtual Size dari sebuah seksi (.text) jauh lebih besar daripada Size of Raw Data pada disk?',
            'options': [
                'File tersebut dikompilasi dengan opsi debug aktif',
                'Seksi tersebut di-pack (packing) atau didekompresi ke memori saat runtime',
                'File tersebut merupakan executable 64-bit murni',
                'Aplikasi menggunakan arsitektur microservices'
            ],
            'correct': 1,
            'explanation': 'Perbedaan drastis di mana Virtual Size jauh melampaui Raw Data adalah karakteristik klasik dari packed malware (seperti UPX/Themida), di mana payload terkompresi kecil di disk akan mengembang di memori saat dieksekusi.'
        },
        {
            'question': 'Mengapa Shannon Entropy sebesar 7.85 bits pada seksi .text tidak boleh langsung disimpulkan sebagai malware tanpa analisis lanjutan?',
            'options': [
                'Karena nilai entropi maksimum adalah 16 bits',
                'Karena perangkat lunak legal, installer, pustaka kriptografi, dan file kompresi juga memiliki entropi mendekati 8.0 bits',
                'Karena entropi hanya berlaku untuk transmisi jaringan nirkabel',
                'Karena model Machine Learning tidak dapat membaca nilai desimal'
            ],
            'correct': 1,
            'explanation': 'Slide presentasi menegaskan: "Feature != truth". Entropi tinggi hanya mengindikasikan ketiadaan struktur byte repetitif (seperti kompresi atau enkripsi). Tanpa konteks sertifikat digital dan perilaku API, kesimpulan malware akan memicu tingginya False Positive.'
        },
        {
            'question': 'Urutan pemanggilan API manakah di bawah ini yang merepresentasikan pola serangan Process Injection (injeksi kode ke proses sah)?',
            'options': [
                'CreateWindow -> ShowWindow -> UpdateWindow -> GetMessage',
                'OpenProcess -> VirtualAllocEx -> WriteProcessMemory -> CreateRemoteThread',
                'WSAStartup -> socket -> connect -> send',
                'RegOpenKey -> RegQueryValue -> RegCloseKey -> ExitProcess'
            ],
            'correct': 1,
            'explanation': 'Ini adalah signature behavioral sekuensial klasik: malware membuka handle proses lain (OpenProcess), memesan alokasi memori remote (VirtualAllocEx), menyuntikkan payload (WriteProcessMemory), lalu mengeksekusinya (CreateRemoteThread).'
        },
        {
            'question': 'Apa kelemahan utama model pendeteksi berbasis Bag-of-API Calls jika dibandingkan dengan pemodelan N-Gram atau RNN/LSTM?',
            'options': [
                'Bag-of-Calls membutuhkan komputasi GPU yang terlalu mahal',
                'Bag-of-Calls kehilangan informasi urutan (order) dan konteks hubungan kronologis antar panggilan fungsi',
                'Bag-of-Calls tidak dapat menghitung frekuensi kemunculan',
                'Bag-of-Calls hanya bekerja pada sistem operasi Linux'
            ],
            'correct': 1,
            'explanation': 'Pada slide ditegaskan: "Frequency loses order; order can encode intent". Bag-of-calls hanya menghitung jumlah, sedangkan niat jahat (seperti urutan injeksi kode) terenkode kuat pada urutan eksekusi sekuensial.'
        },
        {
            'question': 'Pada tumpukan sistem komputer (Computer-System Stack), dimanakah letak kontrol hak istimewa (privilege) tertinggi dan pemantauan paling sulit?',
            'options': [
                'Application Layer (Scripts & User-space modules)',
                'Win32 API Runtime Interface',
                'Kernel & Drivers / Firmware (Ring 0 & Bawah)',
                'Web Browser Extension'
            ],
            'correct': 2,
            'explanation': 'Prinsip utama slide: "The lower the layer, the stronger the control—and the harder the monitoring". Rootkit yang beroperasi pada level Kernel atau Firmware memiliki kontrol tak terbatas dan dapat menyembunyikan diri dari EDR/antivirus user-space.'
        },
        {
            'question': 'Dalam evaluasi model deteksi malware pada lingkungan produksi Enterprise SOC, mengapa metrik Akurasi Standar (Accuracy) sering dihindari?',
            'options': [
                'Karena akurasi membutuhkan waktu pelatihan berkali-kali lipat',
                'Karena rasio file benign dan malware sangat tidak seimbang (class imbalance), sehingga False Positive Rate (FPR) lebih esensial',
                'Karena akurasi hanya dapat diukur untuk file dengan ukuran di bawah 1 MB',
                'Karena analis SOC tidak memahami nilai persentase'
            ],
            'correct': 1,
            'explanation': 'Di dunia nyata, jutaan event benign terjadi setiap hari dibanding segelintir ancaman malware. False Positive Rate yang sedikit saja tinggi akan membanjiri analis SOC dengan alert fatigue.'
        },
        {
            'question': 'Apa peran utama teknik Imphash (Import Hash) dalam melacak rumpun malware yang bermutasi?',
            'options': [
                'Menghasilkan nilai hash MD5 dari urutan nama fungsi API yang diimpor oleh executable untuk melacak kesamaan codebase',
                'Menghapus header PE agar file tidak dapat dianalisis oleh Ghidra atau IDA Pro',
                'Mengukur kapasitas RAM yang digunakan selama proses enkripsi ransomware',
                'Mematikan koneksi internet target secara permanen'
            ],
            'correct': 0,
            'explanation': 'Imphash mengonversi urutan fungsi API impor menjadi hash MD5, memungkinkan analis melacak rumpun (malware family) yang sama meskipun biner telah di-recompile atau diubah hash SHA256-nya.'
        },
        {
            'question': 'Berapakah nilai maksimum teoritis dari Shannon Entropy pada distribusi byte file biner?',
            'options': [
                '4.0 bits',
                '8.0 bits (karena 1 byte = 8 bit, dengan 256 kemungkinan nilai)',
                '16.0 bits',
                '100.0 bits'
            ],
            'correct': 1,
            'explanation': 'Karena 1 byte terdiri dari 8 bit (2^8 = 256 nilai unik 0 s/d 255), entropi Shannon maksimum H(X) adalah -log2(1/256) = 8.0 bits per byte.'
        }
    ]

    selected_quiz = random.sample(quiz_pool, min(6, len(quiz_pool)))
    shuffled_quiz = shuffle_quiz_options(selected_quiz)

    return {
        'title': doc_title if doc_title else 'IF640-W05: The Physics of Malware',
        'slides_count': total_slides,
        'summary': {
            'executive_summary': (
                f"Materi kuliah '{doc_title}' ({total_slides} slides) mengupas transisi fundamental deteksi malware: "
                "dari pendekatan reaktif berbasis signature hash statis (MD5/SHA256) menuju pemodelan perilaku dinamis dan "
                "rekayasa fitur kecerdasan buatan (Machine Learning Pipeline). "
                "Fokus materi mencakup anatomi file PE (Portable Executable), kuantifikasi keacakan byte via Shannon Entropy, "
                "telemetri runtime pada antarmuka OS (API Calls vs System Calls), rantai eksekusi Process Injection, "
                "hingga ekstraksi pola sekuensial menggunakan N-Gram (Bigram & Trigram) untuk melatih model Random Forest, SVM, dan LSTM."
            ),
            'key_takeaways': [
                'Siklus Hidup Eksekusi Malware (Execution Lifecycle): 6 tahapan terstruktur: Delivery -> Load -> Execute -> Persist -> Communicate -> Impact. Deteksi di fase awal (Delivery/Load) mencegah kompromi sistemik.',
                'Anatomi Portable Executable (PE): Header memuat metadata arsitektur, ImageBase, dan AddressOfEntryPoint. Seksi memuat kode biner (.text), variabel global (.data), serta daftar dependensi pustaka (.rdata & Import Address Table). Indikator packing terlihat saat VirtualSize jauh melampaui SizeOfRawData.',
                'Shannon Entropy H(X) = -Σ p(x) log2 p(x): Mengukur keacakan distribusi byte (0.0 s/d 8.0 bits). Nilai di atas 7.2 bits mengindikasikan payload terenkripsi atau packed binary. PENTING: File legal (ZIP, MP4, installer) juga berentropi tinggi, sehingga entropi bukan bukti mutlak malware (feature != truth).',
                'API Calls vs System Calls: Win32 API berada di user-space dengan arti semantik jelas (misal: CreateProcess, VirtualAlloc), sedangkan System Calls adalah transisi gerbang masuk langsung ke Kernel (Ring 0) dengan volume sangat tinggi dan arsitektur spesifik.',
                'Pentingnya Sekuensial (Sequence Matters): Bag-of-Calls kehilangan kronologi. Rangkaian "OpenProcess -> VirtualAllocEx -> WriteProcessMemory -> CreateRemoteThread" adalah fingerprint kuat Process Injection yang dapat ditangkap akurat melalui pemodelan N-Gram (Bigram/Trigram) dan RNN/LSTM.',
                'Evaluasi Model & Adversarial ML: Dalam keamanan siber, akurasi tinggi sering kali menyesatkan akibat ketidakseimbangan kelas (class imbalance). Metrik kritis adalah False Positive Rate (FPR) pada batas deteksi tinggi serta ketahanan terhadap Adversarial Perturbation (seperti dead code insertion & section padding).'
            ],
            'cheat_sheet': [
                {
                    'layer': 'PE Structure & Headers',
                    'unit': 'Static Biner',
                    'protocols': 'AddressOfEntryPoint, Imphash, Section VirtualSize vs RawData',
                    'hardware': 'Mendeteksi packing (UPX/Themida) & silsilah malware family'
                },
                {
                    'layer': 'Shannon Entropy H(X)',
                    'unit': 'Metrik Acak',
                    'protocols': 'H(X) = -Σ p(x) log2 p(x) | Skala 0.0 - 8.0 bits',
                    'hardware': 'Kepadatan byte tinggi (>7.2) = Payload encrypted/compressed'
                },
                {
                    'layer': 'API vs System Calls',
                    'unit': 'Telemetri OS',
                    'protocols': 'User-Space (Win32 semantik) vs Kernel Ring 0 (Syscalls masif)',
                    'hardware': 'Logging via API hooking, EDR sensor, ETW (Event Tracing)'
                },
                {
                    'layer': 'Process Injection Chain',
                    'unit': 'Runtime Pattern',
                    'protocols': 'OpenProcess -> VirtualAllocEx -> WriteProcessMemory -> CreateRemoteThread',
                    'hardware': 'Menyuntikkan shellcode ke proses sah (svchost, explorer)'
                },
                {
                    'layer': 'N-Gram Sequence ML',
                    'unit': 'Feature Vector',
                    'protocols': 'Bigram (n=2), Trigram (n=3) via CountVectorizer / TF-IDF',
                    'hardware': 'Klasifikasi intent sekuensial dengan Random Forest / LSTM'
                },
                {
                    'layer': 'Persistence & Defense',
                    'unit': 'System State',
                    'protocols': 'Run Keys, Scheduled Tasks, Win Services, LOLBins execution',
                    'hardware': 'Evaluasi model ML: Fokus AUC-ROC & minimalkan False Positive'
                }
            ],
            'exam_tips': [
                'Jebakan Ujian Nilai Entropi: Jika soal menanyakan "Apakah file biner dengan entropi 7.9 bit pasti merupakan trojan malware?", jawabannya adalah SALAH! File sah seperti installer game, dokumen PDF terenkripsi, file ZIP, dan video kompresi MP4 juga berentropi tinggi. Harus dikombinasikan dengan korelasi seksi PE dan API import.',
                'Hafalkan Urutan 4 API Call Process Injection: OpenProcess (buka handle proses korban) -> VirtualAllocEx (alokasikan memori di proses target) -> WriteProcessMemory (tulis payload/shellcode ke memori remote) -> CreateRemoteThread (jalankan thread baru untuk eksekusi payload).',
                'Kelemahan Bag-of-Calls: Ingat kata kunci "Frequency loses order, order encodes intent". Bag-of-calls hanya menghitung frekuensi tanpa memedulikan urutan waktu, sehingga gagal mendeteksi pola serangan yang bergantung pada urutan kronologis.',
                'Metrik Evaluasi: Jangan terjebak metrik akurasi (accuracy) semata. Dalam data malware yang sangat imbalanced, model yang memprediksi semua file sebagai "aman" bisa meraih akurasi 99%, namun fatal karena meloloskan ancaman kritis. Gunakan FPR at high recall dan PR-AUC.'
            ]
        },
        'quiz': shuffled_quiz,
        'flashcards': [
            {
                'front': 'Shannon Entropy H(X)',
                'back': 'Rumus: -Σ p(x) log2 p(x). Mengukur keacakan data biner. Skala 0.0 - 8.0 bits. Digunakan untuk mendeteksi enkripsi payload atau packer biner.',
                'badge': 'Matematika Deteksi'
            },
            {
                'front': 'PE Section .text vs .rdata',
                'back': '.text menyimpan instruksi kode mesin (machine code) yang dieksekusi CPU. .rdata menyimpan data read-only dan Import Address Table (IAT).',
                'badge': 'Struktur PE'
            },
            {
                'front': 'Process Injection Sequence',
                'back': 'OpenProcess -> VirtualAllocEx -> WriteProcessMemory -> CreateRemoteThread. Menyuntikkan shellcode berbahaya ke dalam memori proses sah (misal explorer.exe).',
                'badge': 'Behavioral Attack'
            },
            {
                'front': 'Static vs Dynamic Analysis',
                'back': 'Static: Menganalisis PE header, string, dan hash tanpa mengeksekusi biner. Dynamic: Menjalankan sampel di sandbox terisolasi untuk merekam proses, registry, & network.',
                'badge': 'Metodologi Analisis'
            },
            {
                'front': 'Imphash (Import Hash)',
                'back': 'Hash MD5 yang dihitung dari urutan nama fungsi API yang diimpor oleh executable. Berfungsi melacak rumpun (family) malware meskipun file di-recompile.',
                'badge': 'Feature Engineering'
            },
            {
                'front': 'System Calls vs API Calls',
                'back': 'API Calls (Win32) berjalan di user-space dengan arti semantik jelas. System Calls adalah transisi ke Kernel space (Ring 0), bervolume tinggi, dan spesifik versi OS.',
                'badge': 'Arsitektur OS'
            },
            {
                'front': 'N-Gram pada API Sequences',
                'back': 'Memecah deretan pemanggilan API menjadi pasangan n token berurutan (Bigram n=2, Trigram n=3) agar model ML dapat menangkap konteks kronologi niat serangan.',
                'badge': 'Machine Learning'
            },
            {
                'front': 'Malware Execution Lifecycle',
                'back': 'Urutan 6 tahap: 1. Delivery -> 2. Load -> 3. Execute -> 4. Persist -> 5. Communicate -> 6. Impact. Semakin awal dideteksi, semakin minim dampak kerusakannya.',
                'badge': 'Siklus Serangan'
            }
        ],
        'video': {
            'duration': '4:45',
            'lecturer': 'Dr. Maya AI - Virtual Lecturer',
            'chapters': [
                {'time': '00:00', 'title': 'Bab 1: Roadmap & Anatomi Portable Executable (PE)'},
                {'time': '01:10', 'title': 'Bab 2: Shannon Entropy & Indikator Packer'},
                {'time': '02:25', 'title': 'Bab 3: API Calls vs System Calls & Telemetri'},
                {'time': '03:30', 'title': 'Bab 4: Sekuensial N-Gram & Pipeline ML Detection'},
                {'time': '04:10', 'title': 'Bab 5: Ringkasan Ujian & Pertahanan EDR'}
            ],
            'slides': [
                {
                    'slide_num': 1,
                    'title': 'Anatomi Portable Executable (PE)',
                    'subtitle': 'Struktur Biner Windows: Headers, Sections, & Imports',
                    'bullet_points': [
                        'Headers: AddressOfEntryPoint, Machine Architecture (x86/x64), ImageBase.',
                        'Seksi Biner: .text (Executable Code), .data (Globals), .rdata (Import Directory).',
                        'Analisis Imphash: Melacak silsilah malware berdasarkan pola pemanggilan library DLL.'
                    ],
                    'transcript': 'Selamat datang di kuliah The Physics of Malware! Di modul ini, kita membedah bahwa sebuah file biner bukanlah sekadar blok hitam biner, melainkan struktur data PE yang memuat cetak biru lengkap bagaimana malware berniat beroperasi di memori Windows.'
                },
                {
                    'slide_num': 2,
                    'title': 'Fisika Entropi Shannon dalam Deteksi',
                    'subtitle': 'H(X) = -Σ p(x) log2 p(x) — Kepadatan Keacakan Biner',
                    'bullet_points': [
                        'Nilai mendekati 8.0 bits menandakan kompresi tinggi atau enkripsi payload.',
                        'Korelasi krusial: Membandingkan entropi antar seksi biner untuk mendeteksi packing.',
                        'Catatan penting: Jangan langsung menyimpulkan malware hanya dari angka entropi tinggi!'
                    ],
                    'transcript': 'Entropi Shannon mengukur ketidakteraturan byte data. Bila sebuah executable memiliki seksi .text dengan entropi sangat tinggi di atas 7.5 bits, besar kemungkinan kode aslinya dienkripsi atau dipadatkan oleh software packer untuk mengelabui deteksi antivirus statis.'
                },
                {
                    'slide_num': 3,
                    'title': 'API Calls, System Calls, & Process Injection',
                    'subtitle': 'Observasi Perilaku Runtime: User Space vs Kernel Space',
                    'bullet_points': [
                        'API Calls (Win32 API): Memberikan konteks semantik tinggi (OpenFile, CreateProcess).',
                        'System Calls: Gerbang transisi ke Kernel Ring 0 dengan volume data sangat masif.',
                        'Rantai Process Injection: OpenProcess -> VirtualAllocEx -> WriteProcessMemory -> CreateRemoteThread.'
                    ],
                    'transcript': 'Ketika malware berjalan, ia berinteraksi dengan kernel melalui System Calls. Pola berbahaya dapat diidentifikasi secara unik ketika sebuah proses yang tidak dikenal memesan alokasi memori ke proses sistem seperti svchost.exe.'
                },
                {
                    'slide_num': 4,
                    'title': 'Sekuensial N-Gram & Pipeline AI',
                    'subtitle': 'Mengubah Rentetan Event Menjadi Feature Vector ML',
                    'bullet_points': [
                        'Bag-of-calls kehilangan kronologi; pemodelan N-Gram mempertahankan konteks urutan.',
                        'Bigram & Trigram diekstrak dengan CountVectorizer untuk klasifikasi linear / pohon.',
                        'Arsitektur lanjutan: RNN, LSTM, dan Transformer untuk analisis dependensi jangka panjang.'
                    ],
                    'transcript': 'Di era AI modern, deteksi malware tidak lagi mengandalkan signature statis semata. Kita mengekstrak n-gram dari urutan API calls, lalu melatih model machine learning seperti Random Forest atau LSTM untuk membedakan software jinak dari malware canggih.'
                }
            ]
        }
    }

# ==============================================================================
# WEEK 03: CLASSIFY THREATS WITH SUPERVISED LEARNING
# ==============================================================================
def generate_supervised_w03(doc_title, total_slides):
    quiz_pool = [
        {
            'question': 'Apa perbedaan mendasar antara klasifikasi Binary vs Multiclass dalam deteksi ancaman keamanan siber?',
            'options': [
                'Binary memisahkan sampel menjadi Benign vs Malicious, sedangkan Multiclass mengklasifikasikan sampel ke dalam keluarga malware spesifik (misal: Ransomware, Spyware, Trojan)',
                'Binary hanya dapat dijalankan pada CPU 2-core',
                'Multiclass tidak memerlukan label ground truth',
                'Binary hanya memproses file biner berekstensi .bin'
            ],
            'correct': 0,
            'explanation': 'Binary classification menjawab pertanyaan "apakah file ini berbahaya atau aman?", sedangkan Multiclass classification mengidentifikasi tipe atau rumpun keluarga malware spesifik.'
        },
        {
            'question': 'Dalam evaluasi model deteksi malware pada data yang sangat tidak seimbang (class imbalance), mengapa metrik Precision dan Recall lebih diutamakan daripada Accuracy?',
            'options': [
                'Karena pada data imbalance (misal 99.9% aman, 0.1% malware), model yang menebak semua aman meraih akurasi 99.9% tetapi gagal mendeteksi ancaman sama sekali',
                'Karena Precision dan Recall tidak membutuhkan pembagian data testing',
                'Karena Accuracy hanya berlaku untuk model Neural Network',
                'Karena analis SOC tidak dapat menghitung persentase'
            ],
            'correct': 0,
            'explanation': 'Accuracy Paradox terjadi saat mayoritas data benign membuat akurasi tampak tinggi semu. Precision mengukur minimnya false positive, sedangkan Recall mengukur seberapa sedikit malware yang lolos (false negative).'
        },
        {
            'question': 'Bagaimana cara kerja algoritma Support Vector Machine (SVM) dalam memisahkan data serangan dari lalu lintas jaringan normal?',
            'options': [
                'Mencari hyperplane pemisah optimal yang memaksimalkan margin jarak antara support vectors dari kedua kelas data',
                'Menghitung nilai rata-rata dari seluruh baris tabel secara acak',
                'Mengurutkan data berdasarkan timestamp kedatangan paket',
                'Memecah file biner menjadi pecahan 1 KB'
            ],
            'correct': 0,
            'explanation': 'SVM bekerja dengan mengidentifikasi support vectors (titik terdekat ke batas kelas) dan menemukan hyperplane yang memaksimalkan margin pemisah antar kelas.'
        }
    ]
    shuffled_quiz = shuffle_quiz_options(quiz_pool)

    return {
        'title': doc_title if doc_title else 'IF640-W03: Classify Threats with Supervised Learning',
        'slides_count': total_slides,
        'summary': {
            'executive_summary': (
                f"Materi kuliah '{doc_title}' ({total_slides} slides) mengupas tuntas penerapan Supervised Machine Learning "
                "dalam klasifikasi ancaman siber. Pembahasan menitikberatkan pada pembedaan fitur vs label, alur kerja pipeline "
                "pelatihan (Train, Validation, Test), algoritma Decision Trees, Random Forest, dan Support Vector Machine (SVM), "
                "serta mitigasi class imbalance menggunakan metrik evaluasi Confusion Matrix, Precision, Recall, dan PR-AUC."
            ),
            'key_takeaways': [
                'Supervised Learning Pipeline: Membutuhkan data latih berlabel ground-truth. Terbagi menjadi fase feature extraction, data splitting (Train/Val/Test), training, dan evaluasi.',
                'Binary vs Multiclass Classification: Binary memisahkan benign vs malicious; multiclass mengelompokkan ke dalam kategori serangan spesifik (misal: Trojan, Ransomware, Adware).',
                'Decision Trees & Ensembles: Decision Tree mudah diinterpretasikan oleh analis; Random Forest dan Gradient Boosting mengurangi variance dan overfitting dengan menggabungkan banyak pohon.',
                'Support Vector Machines (SVM): Memaksimalkan margin pemisah antar kelas menggunakan kernel trick untuk data non-linear.',
                'Mitigasi Class Imbalance: Dalam dataset keamanan riil, serangan malware jauh lebih langka daripada trafik benign. Evaluasi harus mengandalkan PR-AUC, F1-Score, dan Recall, bukan Accuracy semata.'
            ],
            'cheat_sheet': [
                {'layer': 'Data Splitting', 'unit': 'Pipeline', 'protocols': 'Train (60-70%), Validation (15-20%), Test (15-20%)', 'hardware': 'Mencegah data leakage dan overfitting'},
                {'layer': 'Decision Tree', 'unit': 'Tree Model', 'protocols': 'Gini Impurity, Information Gain (Entropy)', 'hardware': 'Explainable rule-based decision paths'},
                {'layer': 'Random Forest', 'unit' : 'Ensemble', 'protocols': 'Bagging + Random Subspace Sampling', 'hardware': 'High accuracy baseline, robust to noise'},
                {'layer': 'SVM Classifier', 'unit': 'Max-Margin', 'protocols': 'Linear, RBF Kernel, Support Vectors', 'hardware': 'Optimal hyperplane boundary in high-dim space'},
                {'layer': 'Evaluation Metrics', 'unit': 'SOC Metrics', 'protocols': 'Confusion Matrix: TP, FP, TN, FN | Precision, Recall', 'hardware': 'Prioritas: Minimalisir False Positive & False Negative'}
            ],
            'exam_tips': [
                'Ingat paradoks akurasi: Pada dataset SOC dengan 99.9% benign dan 0.1% malware, akurasi 99.9% tidak bernilai jika recall malware adalah 0%.',
                'Pahami rumus Precision = TP / (TP + FP) dan Recall = TP / (TP + FN).'
            ]
        },
        'quiz': shuffled_quiz,
        'flashcards': [
            {'front': 'Features vs Labels', 'back': 'Features: variabel input terukur (misal: packet size, port, entropy). Labels: target ground-truth (misal: benign=0, malware=1).', 'badge': 'Supervised ML'},
            {'front': 'Confusion Matrix', 'back': 'Tabel 2x2: True Positive (TP), False Positive (FP), True Negative (TN), False Negative (FN).', 'badge': 'Evaluasi'},
            {'front': 'Random Forest', 'back': 'Ensemble learning yang menggabungkan prediksi ratusan Decision Trees via voting mayoritas untuk menekan varians.', 'badge': 'Algoritma'},
            {'front': 'Support Vector Machine (SVM)', 'back': 'Algoritma klasifikasi yang mencari hyperplane dengan margin terjauh antara dua kelas data.', 'badge': 'Algoritma'}
        ],
        'video': {
            'duration': '4:10',
            'lecturer': 'Dr. Maya AI - Virtual Lecturer',
            'chapters': [
                {'time': '00:00', 'title': 'Bab 1: Konsep Supervised Learning dalam Keamanan Siber'},
                {'time': '01:15', 'title': 'Bab 2: Decision Trees, Random Forest, & SVM'},
                {'time': '02:40', 'title': 'Bab 3: Evaluasi Model & Mengatasi Class Imbalance'}
            ],
            'slides': [
                {
                    'slide_num': 1,
                    'title': 'Klasifikasi Ancaman Supervised Learning',
                    'subtitle': 'Memetakan Pola Telemetri ke Label Ground-Truth',
                    'bullet_points': ['Pelatihan model dengan dataset teranotasi.', 'Pembedaan binary benign vs malicious.', 'Multiclass malware family classification.'],
                    'transcript': 'Selamat datang di modul Supervised Learning untuk Cybersecurity! Kita akan mempelajari bagaimana melatih model pohon dan SVM untuk mengenali tanda-tanda serangan siber secara otomatis.'
                }
            ]
        }
    }

# ==============================================================================
# WEEK 02: HOW AI THINKS ABOUT DATA
# ==============================================================================
def generate_data_w02(doc_title, total_slides):
    quiz_pool = [
        {
            'question': 'Apa perbedaan mendasar antara Structured Data dan Unstructured Data dalam telemetri keamanan siber?',
            'options': [
                'Structured data memiliki skema baris-kolom terstandarisasi (seperti database relational atau netflow tabel), sedangkan Unstructured data tidak memiliki format kaku (seperti raw packet payload atau teks laporan ancaman)',
                'Structured data tidak dapat disimpan di komputer',
                'Unstructured data hanya diproses menggunakan kabel LAN',
                'Structured data tidak pernah memuat alamat IP'
            ],
            'correct': 0,
            'explanation': 'Structured data memiliki tipe data dan skema terdefinisi (kolom IP, port, bytes), sedangkan Unstructured data memerlukan teknik ekstraksi teks/fitur khusus sebelum diproses model ML.'
        }
    ]
    shuffled_quiz = shuffle_quiz_options(quiz_pool)

    return {
        'title': doc_title if doc_title else 'IF640-W02: How AI Thinks about Data',
        'slides_count': total_slides,
        'summary': {
            'executive_summary': (
                f"Materi kuliah '{doc_title}' ({total_slides} slides) membahas fondasi representasi data keamanan siber "
                "dalam pemodelan Artificial Intelligence. Fokus materi meliputi transisi dari pendekatan logic-driven (rule-based) "
                "ke data-driven, pemrosesan structured vs unstructured data (firewall logs, sysmon, network packets), "
                "serta teknik feature engineering, encoding kategorikal, dan normalisasi vektor."
            ),
            'key_takeaways': [
                'Logic-Driven vs Data-Driven: Aturan statis (IF-THEN) mudah dihindari oleh penyerang; model data-driven belajar dari distribusi statistik fitur.',
                'Karakteristik Structured Data: Terdiri dari skema kolom tetap (Timestamp, Source IP, Dest Port, Action) yang siap dimasukkan ke pipeline ML numerik.',
                'Representasi Fitur & Encoding: Fitur kategorikal (protokol TCP/UDP) di-encode menggunakan One-Hot Encoding; fitur teks diekstrak via TF-IDF atau embeddings.',
                'Pembersihan & Sanitasi Log: Menangani missing values, parsing format waktu ISO, dan agregasi window waktu untuk mengurangi noise telemetri.'
            ],
            'cheat_sheet': [
                {'layer': 'Structured Data', 'unit': 'Tabel / CSV', 'protocols': 'NetFlow, Firewall Logs, Auth Records', 'hardware': 'Kolom terdefinisi, siap diproses ML'},
                {'layer': 'Unstructured Data', 'unit': 'Biner / Teks', 'protocols': 'PCAP payloads, Memory Dumps, Threat Reports', 'hardware': 'Membutuhkan tokenisasi & ekstraksi embedding'},
                {'layer': 'Feature Scaling', 'unit': 'Vektor', 'protocols': 'MinMaxScaler [0,1], StandardScaler (Z-score)', 'hardware': 'Menjaga kestabilan komputasi algoritma berbasis jarak'}
            ],
            'exam_tips': ['Pahami perbedaan tipe data kategorikal, numerik diskrit, dan kontinu dalam log firewall.']
        },
        'quiz': shuffled_quiz,
        'flashcards': [
            {'front': 'Structured vs Unstructured Data', 'back': 'Structured: tabular dengan skema tetap (log database). Unstructured: teks bebas, payload biner pcap, laporan PDF.', 'badge': 'Data Engineering'},
            {'front': 'One-Hot Encoding', 'back': 'Mengubah variabel kategorikal (misal: TCP, UDP, ICMP) menjadi vektor biner bernilai 0 dan 1.', 'badge': 'Encoding'}
        ],
        'video': {
            'duration': '3:50',
            'lecturer': 'Dr. Maya AI - Virtual Lecturer',
            'chapters': [{'time': '00:00', 'title': 'Bab 1: Representasi Data Keamanan Siber untuk AI'}],
            'slides': [{'slide_num': 1, 'title': doc_title, 'subtitle': 'Dari Log Mentah Menjadi Vektor Fitur', 'bullet_points': ['Struktur data terorganisir', 'Normalisasi data log'], 'transcript': 'Selamat datang di pembahasan How AI Thinks About Data!'}]
        }
    }

# ==============================================================================
# WEEK 01: THE AI CYBER LANDSCAPE
# ==============================================================================
def generate_landscape_w01(doc_title, total_slides):
    quiz_pool = [
        {
            'question': 'Apa faktor utama yang memicu perlunya adopsi AI dalam Security Operations Center (SOC) modern?',
            'options': [
                'Volume data log dan event keamanan yang melampaui kemampuan analisis manual manusia serta maraknya serangan otomatis berkecepatan tinggi',
                'Kewajiban menghapus semua analis manusia dari SOC',
                'AI dapat mencegah mati lampu pada datacenter',
                'Antivirus tradisional tidak lagi memerlukan koneksi listrik'
            ],
            'correct': 0,
            'explanation': 'Eksplosi volume telemetri (jutaan log per hari) dan alert fatigue menuntut sistem cerdas berbasis AI untuk memfilter, mengkorelasikan, dan merespons ancaman secara real-time.'
        }
    ]
    shuffled_quiz = shuffle_quiz_options(quiz_pool)

    return {
        'title': doc_title if doc_title else 'IF640-W01: The AI Cyber Landscape',
        'slides_count': total_slides,
        'summary': {
            'executive_summary': (
                f"Materi kuliah '{doc_title}' ({total_slides} slides) memberikan tinjauan strategis evolusi lanskap ancaman siber "
                "dan peran Artificial Intelligence dalam pertahanan modern. Pembahasan mencakup sejarah evolusi ancaman, "
                "analisis kasus pelanggaran data historis (Target & Sony attacks), perluasan permukaan serangan (Attack Surface), "
                "serta paradigma dual-use AI (Defensive AI vs Adversarial AI)."
            ),
            'key_takeaways': [
                'Evolusi Lanskap Ancaman: Serangan telah bertransformasi dari virus berbasis signature sederhana menjadi kampanye APT (Advanced Persistent Threat) yang canggih dan terkoordinasi.',
                'Perluasan Permukaan Serangan (Attack Surface): Integrasi cloud, IoT, dan ekosistem SaaS memperbesar vektor masuknya musuh ke dalam infrastruktur organisasi.',
                'Dual-Use AI: AI digunakan oleh defender untuk deteksi anomali cepat, namun juga dimanfaatkan penyerang untuk otomatisasi phishing dan polymorphism payload.',
                'Tantangan SOC Modern: Mengatasi banjir peringatan (Alert Fatigue) melalui otomatisasi triase dan korelasi telemetri cerdas.'
            ],
            'cheat_sheet': [
                {'layer': 'Defensive AI', 'unit': 'Deteksi', 'protocols': 'EDR, SIEM, SOAR, Anomaly Detection', 'hardware': 'Automated triage & rapid incident response'},
                {'layer': 'Offensive AI', 'unit': 'Evasion', 'protocols': 'Deepfake Phishing, Automated Recon, Evasion', 'hardware': 'Serangan berskala masif dengan personalisasi tinggi'},
                {'layer': 'Attack Surface', 'unit': 'Eksposur', 'protocols': 'Cloud Assets, Shadow IT, Endpoints, APIs', 'hardware': 'Pemetaan eksposur sistem terhadap vektor serangan'}
            ],
            'exam_tips': ['Pahami dampak kasus kebocoran data Target dan Sony sebagai tonggak sejarah penting evolusi keamanan siber perusahaan.']
        },
        'quiz': shuffled_quiz,
        'flashcards': [
            {'front': 'Attack Surface', 'back': 'Seluruh titik masuk dan kerentanan yang dapat dieksploitasi oleh adversary untuk mengompromikan sistem.', 'badge': 'Konsep Dasar'},
            {'front': 'Alert Fatigue', 'back': 'Kondisi kelelahan analis keamanan akibat banjir ribuan notifikasi alarm palsu (false positives) setiap harinya.', 'badge': 'SOC Challenge'}
        ],
        'video': {
            'duration': '3:45',
            'lecturer': 'Dr. Maya AI - Virtual Lecturer',
            'chapters': [{'time': '00:00', 'title': 'Bab 1: Evolusi Ancaman Siber & Peran AI'}],
            'slides': [{'slide_num': 1, 'title': doc_title, 'subtitle': 'Lanskap Keamanan Siber di Era Kecerdasan Buatan', 'bullet_points': ['Evolusi ancaman modern', 'Dual-use AI'], 'transcript': 'Selamat datang di kuliah perdana AI in Cybersecurity!'}]
        }
    }

# ==============================================================================
# WEEK 06: SOCIAL ENGINEERING
# ==============================================================================
def generate_social_engineering_w06(doc_title, total_slides):
    quiz_pool = [
        {
            'question': 'Apa vektor manipulasi psikologis yang paling sering dimanfaatkan dalam serangan Business Email Compromise (BEC)?',
            'options': [
                'Otoritas (Authority) dan Urgensi (Urgency) dengan menyamar sebagai pimpinan eksekutif yang meminta transfer dana darurat',
                'Rasa penasaran terhadap diskon belanja online semata',
                'Ajakan bermain game online multiplayer',
                'Pemberitahuan perubahan cuaca lokal'
            ],
            'correct': 0,
            'explanation': 'Penyerang mengeksploitasi kepatuhan karyawan terhadap otoritas pimpinan (CEO/CFO fraud) dikombinasikan dengan tekanan batas waktu (urgency) agar korban tidak sempat memverifikasi prosedur keuangan.'
        }
    ]
    shuffled_quiz = shuffle_quiz_options(quiz_pool)

    return {
        'title': doc_title if doc_title else 'IF640-W06: Social Engineering',
        'slides_count': total_slides,
        'summary': {
            'executive_summary': (
                f"Materi kuliah '{doc_title}' ({total_slides} slides) membedah teknik manipulasi psikologis dalam rekayasa sosial "
                "(Social Engineering). Pembahasan mencakup taksonomi serangan (Phishing, Spear Phishing, Whaling, Vishing, Smishing, Pretexting), "
                "enam prinsip manipulasi psikologis Cialdini, teknik bypass Multi-Factor Authentication (MFA Fatigue), "
                "serta strategi pertahanan teknis (DMARC, DKIM, SPF) dan Security Awareness Training."
            ),
            'key_takeaways': [
                'Definisi Social Engineering: Eksploitasi kerentanan faktor manusia (human factor) untuk memperoleh akses tidak sah atau data rahasia.',
                'Vektor Serangan Utama: Phishing (massal), Spear Phishing (tertuju), Whaling (eksekutif level atas), Pretexting (skenario identitas palsu), Baiting (umpan flashdisk/software).',
                'Prinsip Psikologis Kunci: Otoritas, Urgensi, Kelangkaan (Scarcity), Konsensus Sosial, dan Rasa Takut.',
                'Pertahanan Terintegrasi: Kombinasi filter email teknis (SPF/DKIM/DMARC), implementasi Zero Trust, hardware security keys (FIDO2), serta pelatihan simulasi phishing berkala.'
            ],
            'cheat_sheet': [
                {'layer': 'Phishing Taxonomy', 'unit': 'Email / Web', 'protocols': 'Spear Phishing, Whaling, Clone Phishing', 'hardware': 'Kredensial harvesting & malware delivery'},
                {'layer': 'Voice & Mobile', 'unit': 'Telefoni', 'protocols': 'Vishing (Voice), Smishing (SMS)', 'hardware': 'OTP theft & identity pretexting'},
                {'layer': 'Email Authentication', 'unit': 'DNS Security', 'protocols': 'SPF, DKIM, DMARC alignment', 'hardware': 'Mencegah spoofing domain resmi organisasi'}
            ],
            'exam_tips': ['Pahami peran DMARC dalam mencegah email domain spoofing pada serangan spear phishing.']
        },
        'quiz': shuffled_quiz,
        'flashcards': [
            {'front': 'Spear Phishing vs Whaling', 'back': 'Spear Phishing: serangan terarah pada individu/organisasi spesifik. Whaling: serangan spear phishing yang khusus menargetkan eksekutif tingkat C (CEO, CFO).', 'badge': 'Social Engineering'},
            {'front': 'MFA Fatigue Attack', 'back': 'Serangan membombardir ponsel korban dengan ratusan notifikasi push MFA hingga korban lelah dan tidak sengaja menekan approve.', 'badge': 'Attack Vector'}
        ],
        'video': {
            'duration': '3:50',
            'lecturer': 'Dr. Maya AI - Virtual Lecturer',
            'chapters': [{'time': '00:00', 'title': 'Bab 1: Anatomi Manipulasi Rekayasa Sosial'}],
            'slides': [{'slide_num': 1, 'title': doc_title, 'subtitle': 'Faktor Manusia dalam Keamanan Siber', 'bullet_points': ['Taktik psikologis serangan', 'Pertahanan email SPF/DKIM/DMARC'], 'transcript': 'Selamat datang di modul Social Engineering!'}]
        }
    }

# ==============================================================================
# ADAPTIVE GENERATOR FOR ANY ARBITRARY PRESENTATION
# ==============================================================================
def generate_adaptive_data(doc_title, filename, slides, combined_text):
    total_slides = len(slides) if slides else 15
    
    # Filter out boilerplate lines
    boilerplate_words = [
        'learning outcome', 'elo ', 'clo ', 'subclo', 'review', 'outline', 'agenda',
        'description', 'evaluation', 'assignment', 'tugas', 'references', 'next week',
        'visi', 'misi', 'daftar pustaka', 'summary', 'assessment', 'marlinda', 'overbeek',
        'if640', 'ai in cybersecurity'
    ]
    
    substantive_points = []
    substantive_topics = []
    
    for s in slides:
        for line in s['lines']:
            line_str = line.strip()
            # Clean separators
            cleaned = line_str
            if '·' in cleaned:
                cleaned = cleaned.split('·')[-1].strip()
            if '—' in cleaned and len(cleaned.split('—')[0]) < 25:
                cleaned = cleaned.split('—')[-1].strip()
                
            line_lower = cleaned.lower()
            if len(cleaned) > 12 and not any(b in line_lower for b in boilerplate_words):
                if not cleaned.isdigit() and not cleaned.startswith('===') and len(cleaned) < 120:
                    if cleaned not in substantive_points:
                        substantive_points.append(cleaned)
                    if len(cleaned) < 60 and cleaned not in substantive_topics:
                        substantive_topics.append(cleaned)

    if len(substantive_points) < 6:
        substantive_points = [
            f"Analisis fundamental materi perkuliahan '{doc_title}'",
            "Identifikasi struktur konseptual, relasi data, dan metodologi perancangan sistem",
            "Mekanisme operasional dan arsitektur komponen modul",
            "Proses verifikasi, pengujian performa, dan validasi fungsional",
            "Studi kasus implementasi dan evaluasi batas operasional",
            "Strategi optimalisasi sistem dan pemecahan kendala teknis"
        ]

    topic_1 = substantive_topics[0] if len(substantive_topics) > 0 else substantive_points[0][:40]
    topic_2 = substantive_topics[1] if len(substantive_topics) > 1 else substantive_points[1][:40]
    topic_3 = substantive_topics[2] if len(substantive_topics) > 2 else substantive_points[2][:40]
    topic_4 = substantive_topics[3] if len(substantive_topics) > 3 else substantive_points[3][:40]

    exec_summary = (
        f"Materi perkuliahan '{doc_title}' ({total_slides} slide) telah dianalisis dan dirangkum secara mendalam oleh AI. "
        f"Fokus pembahasan materi meliputi prinsip inti '{topic_1}', arsitektur dan mekanisme kerja '{topic_2}', "
        f"serta penerapan terapan pada '{topic_3}'. Dokumen ini menyajikan poin-poin kunci esensial, "
        "tabel referensi cepat, dan panduan belajar terstruktur untuk persiapan ujian dan kuis akademik."
    )

    takeaways = [
        f"Fondasi & Definisi Utama: {substantive_points[0]}",
        f"Arsitektur & Komponen Sistem: {substantive_points[1] if len(substantive_points)>1 else 'Pemisahan tanggung jawab antar modul dan integrasi interface.'}",
        f"Mekanisme & Alur Operasional: {substantive_points[2] if len(substantive_points)>2 else 'Aliran pemrosesan data terstruktur dari input hingga evaluasi output.'}",
        f"Karakteristik & Spesifikasi Kunci: {substantive_points[3] if len(substantive_points)>3 else 'Parameter pembatas, asumsi kerja, dan standar operasional.'}",
        f"Penerapan & Studi Kasus: {substantive_points[4] if len(substantive_points)>4 else 'Implementasi praktis pada lingkungan terapan.'}",
        f"Evaluasi & Optimasi: {substantive_points[5] if len(substantive_points)>5 else 'Pengukuran metrik efisiensi dan mitigasi kendala operasional.'}"
    ]

    cheat_sheet = [
        {'layer': f'Bab 1: {topic_1[:25]}', 'unit': 'Fondasi Teori', 'protocols': substantive_points[0][:50], 'hardware': 'Prinsip dasar & latar belakang'},
        {'layer': f'Bab 2: {topic_2[:25]}', 'unit': 'Arsitektur Modul', 'protocols': substantive_points[1][:50] if len(substantive_points)>1 else 'Struktur komponen', 'hardware': 'Mekanisme & interaksi data'},
        {'layer': f'Bab 3: {topic_3[:25]}', 'unit': 'Operasional', 'protocols': substantive_points[2][:50] if len(substantive_points)>2 else 'Alur kerja runtime', 'hardware': 'Studi kasus & implementasi'},
        {'layer': f'Bab 4: {topic_4[:25]}', 'unit': 'Evaluasi Teknis', 'protocols': substantive_points[3][:50] if len(substantive_points)>3 else 'Validasi parameter', 'hardware': 'Metrik pengujian & performa'}
    ]

    adaptive_quiz_pool = [
        {
            'question': f'Berdasarkan pemaparan dalam materi "{doc_title}", apa prinsip esensial yang mendasari {topic_1}?',
            'options': [
                substantive_points[0][:90],
                'Pengabaian dokumentasi dan pengujian sistem berkala',
                'Penyimpanan data unencrypted pada jaringan publik tanpa proteksi',
                'Penggunaan dependensi eksternal tanpa validasi parameter'
            ],
            'correct': 0,
            'explanation': f'Slide presentasi menekankan bahwa {substantive_points[0][:80]} merupakan fondasi konseptual utama dalam materi ini.'
        },
        {
            'question': f'Pada bagian pembahasan {topic_2}, manakah pernyataan yang paling akurat mengenai mekanisme operasionalnya?',
            'options': [
                'Komponen dirancang untuk memperlambat alur eksekusi komputasi',
                substantive_points[1][:90] if len(substantive_points)>1 else 'Pemisahan modul terstruktur untuk memastikan skalabilitas dan keandalan sistem',
                'Sistem tidak mendukung pemrosesan data secara multi-user',
                'Seluruh konfigurasi wajib di-hardcode tanpa file environment'
            ],
            'correct': 1,
            'explanation': 'Slide menggarisbawahi bahwa struktur arsitektur modular menjamin interoperabilitas serta efisiensi eksekusi antar modul.'
        },
        {
            'question': f'Apa tujuan utama dari implementasi metodologi yang dipaparkan pada {topic_3}?',
            'options': [
                'Menghapus integritas data audit pengguna',
                'Membatasi fungsionalitas aplikasi hanya pada satu lingkungan lokal',
                substantive_points[2][:90] if len(substantive_points)>2 else 'Meningkatkan akurasi analisis dan meminimalkan tingkat kesalahan (error rate)',
                'Meningkatkan biaya lisensi perangkat keras tanpa dampak performa'
            ],
            'correct': 2,
            'explanation': 'Metodologi yang sistematis dirancang untuk menjaga konsistensi, keandalan, dan efektivitas hasil analisis data.'
        }
    ]

    selected_adaptive_quiz = random.sample(adaptive_quiz_pool, min(len(adaptive_quiz_pool), len(adaptive_quiz_pool)))
    shuffled_adaptive_quiz = shuffle_quiz_options(selected_adaptive_quiz)

    return {
        'title': doc_title,
        'slides_count': total_slides,
        'summary': {
            'executive_summary': exec_summary,
            'key_takeaways': takeaways,
            'cheat_sheet': cheat_sheet,
            'exam_tips': [
                f'Kuasai Definisi Inti: Pahami secara mendalam konsep "{topic_1}".',
                f'Perhatikan Hubungan Komponen: Dosen pengampu umumnya menguji bagaimana "{topic_2}" berinteraksi dengan "{topic_3}".',
                'Cermati perbedaan terminologi teknis pada diagram perbandingan slide.'
            ]
        },
        'quiz': shuffled_adaptive_quiz,
        'flashcards': [
            {'front': f'Definisi {doc_title[:30]}', 'back': substantive_points[0], 'badge': 'Konsep Inti'},
            {'front': f'{topic_1[:30]}', 'back': substantive_points[1] if len(substantive_points)>1 else 'Prinsip teoritis sistem.', 'badge': 'Fondasi'},
            {'front': f'{topic_2[:30]}', 'back': substantive_points[2] if len(substantive_points)>2 else 'Mekanisme interaksi data.', 'badge': 'Arsitektur'}
        ],
        'video': {
            'duration': '3:45',
            'lecturer': 'Dr. Maya AI - Virtual Lecturer',
            'chapters': [{'time': '00:00', 'title': f'Bab 1: Pengantar {doc_title[:25]}'}],
            'slides': [{'slide_num': 1, 'title': doc_title, 'subtitle': 'Ringkasan Materi AI', 'bullet_points': [substantive_points[0][:60]], 'transcript': f'Selamat datang di materi {doc_title}!'}]
        }
    }

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(json.dumps({'error': 'Missing file path argument'}))
        sys.exit(1)
        
    file_path = sys.argv[1]
    refresh_quiz_only = '--refresh-quiz' in sys.argv
    
    gemini_key = ''
    orig_title = ''
    orig_filename = ''
    
    for i, arg in enumerate(sys.argv):
        if arg == '--gemini-key' and i + 1 < len(sys.argv):
            gemini_key = sys.argv[i + 1]
        elif arg == '--title' and i + 1 < len(sys.argv):
            orig_title = sys.argv[i + 1]
        elif arg == '--filename' and i + 1 < len(sys.argv):
            orig_filename = sys.argv[i + 1]

    if not os.path.exists(file_path):
        print(json.dumps({'error': f'File not found: {file_path}'}))
        sys.exit(1)
        
    slides = extract_text_from_file(file_path)
    result = analyze_and_generate(file_path, slides, gemini_key, orig_title, orig_filename, is_refresh_quiz=refresh_quiz_only)
    
    if refresh_quiz_only:
        print(json.dumps({'quiz': result.get('quiz', [])}, ensure_ascii=True))
    else:
        print(json.dumps(result, ensure_ascii=True))
